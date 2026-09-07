# -*- coding: utf-8 -*-
"""
Script de Exportação Profissional do Seminário Mirage para DOCX e PDF.
Gera cópias formatadas dos 3 documentos de gestão do seminário para envio aos integrantes.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from fpdf import FPDF

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "documentos_seminario")
os.makedirs(OUTPUT_DIR, exist_ok=True)

print(f"Diretorio de saida: {OUTPUT_DIR}")

# =========================================================================
# HELPER DE ESTILIZAÇÃO PARA DOCX
# =========================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12 if level==1 else 8)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(20, 50, 110)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(35, 75, 140)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(50, 90, 160)
    return h

# =========================================================================
# HELPER DE PDF COM FPDF2
# =========================================================================
class CustomPDF(FPDF):
    def __init__(self, title_doc):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.title_doc = title_doc
        self.set_auto_page_break(auto=True, margin=15)
        # Carrega fontes Unicode do Windows para suporte completo a acentos e caracteres especiais
        self.add_font("Arial", "", r"C:\Windows\Fonts\arial.ttf")
        self.add_font("Arial", "B", r"C:\Windows\Fonts\arialbd.ttf")
        self.add_font("Arial", "I", r"C:\Windows\Fonts\ariali.ttf")

    def header(self):
        self.set_font("Arial", "B", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 6, "PROJETO MIRAGE - UNISENAI 2026 | INTELIGÊNCIA ARTIFICIAL", border=0, align="L")
        self.cell(0, 6, self.title_doc, border=0, align="R")
        self.ln(8)
        self.set_draw_color(200, 200, 200)
        self.line(10, 15, 200, 15)
        self.ln(3)

    def footer(self):
        self.set_y(-12)
        self.set_font("Arial", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, f"Página {self.page_no()}/{{nb}}", align="C")

    def add_title_block(self, main_title, subtitle):
        self.add_page()
        self.set_font("Arial", "B", 18)
        self.set_text_color(20, 50, 110)
        self.multi_cell(0, 8, main_title, align="C")
        self.ln(2)
        self.set_font("Arial", "I", 11)
        self.set_text_color(80, 80, 80)
        self.multi_cell(0, 6, subtitle, align="C")
        self.ln(6)
        self.set_draw_color(30, 80, 160)
        self.set_line_width(0.8)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(6)

    def section_heading(self, title):
        self.ln(4)
        self.set_font("Arial", "B", 13)
        self.set_text_color(20, 50, 120)
        self.cell(0, 7, title, fill=False)
        self.ln(8)

    def sub_heading(self, title):
        self.ln(2)
        self.set_font("Arial", "B", 10.5)
        self.set_text_color(40, 80, 140)
        self.cell(0, 6, title)
        self.ln(6)

    def body_text(self, text, bold_prefix=""):
        self.set_font("Arial", "", 9.5)
        self.set_text_color(30, 30, 30)
        if bold_prefix:
            self.set_font("Arial", "B", 9.5)
            self.write(5, bold_prefix + " ")
            self.set_font("Arial", "", 9.5)
        self.write(5, text)
        self.ln(6)

    def bullet_item(self, bold_part, text_part):
        self.set_font("Arial", "B", 9.5)
        self.set_text_color(20, 50, 110)
        self.write(5, f"- {bold_part} ")
        self.set_font("Arial", "", 9.5)
        self.set_text_color(40, 40, 40)
        self.write(5, text_part)
        self.ln(5.5)

# =========================================================================
# 1. GERAR: DIVISÃO DE TAREFAS & APRESENTADORES
# =========================================================================
def gerar_divisao_tarefas():
    print("Gerando Divisao de Tarefas (DOCX & PDF)...")
    
    # DOCX
    doc = Document()
    add_styled_heading(doc, "👥 Divisão de Atividades & Apresentadores — Projeto Mirage", level=1)
    
    p = doc.add_paragraph()
    format_paragraph(p)
    p.add_run("Disciplina: Inteligência Artificial — Engenharia de Controle e Automação — UNISENAI\nOrientador: Me. Ricardo Martinez Vicentini\nTempo Total: 15 Minutos (14 min apresentação + 1 min FAQ)").italic = True
    
    add_styled_heading(doc, "📌 Atribuições Individuais e Temas dos Blocos", level=2)
    
    integrantes = [
        ("Murilo Lameira", "Apresentador 1", "0:00 às 3:00 (3.0 min)", "Slides 1 a 3",
         "Abertura oficial, contextualização do problema de NPCs determinísticos clássicos (FSMs e Behavior Trees previsíveis), motivação para IA evolutiva e visão geral da proposta Mirage."),
        ("Leonardo Retori", "Apresentador 2", "3:00 às 6:30 (3.5 min)", "Slides 4 a 6",
         "Fundamentação de Qualidade-Diversidade (MAP-Elites - Kirk & Scirea), Heatmap 3x3 dos nichos fenotípicos, catálogo SEC de Glavin & Madden (DDA) e Operadores Genéticos (Torneio k=3, Crossover Uniforme, Mutação Gaussiana/EDS)."),
        ("Henry Matheus", "Apresentador 3", "6:30 às 10:00 (3.5 min)", "Slides 7 a 9",
         "Cinemática física da esquiva preditiva (Steering Behaviors de Craig Reynolds e CPA - Ponto de Maior Aproximação), vetor de 4 genes, Orçamento Global de Atributos (B = 1.8) e Função de Fitness Multi-Objetivo com Fator de Mérito."),
        ("Murilo Romualdo", "Apresentador 4", "10:00 às 14:00 (4.0 min)", "Slides 10 a 14",
         "Arena dinâmica 2D com 4 pilares de cobertura física e padrões de Bullet Hell (leque e espiral), mitigação de Noisy Fitness com curvas monotônicas, Validação Estatística One-Way ANOVA (F = 138.24, p = 1.44e-15) e Boxplots, Demonstração em GIF e Conclusões.")
    ]
    
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Integrante"
    hdr[1].text = "Papel"
    hdr[2].text = "Tempo"
    hdr[3].text = "Slides"
    hdr[4].text = "Foco Técnico e Responsabilidades"
    for c in hdr:
        set_cell_background(c, "204080")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)
                
    for nome, papel, tempo, slides, desc in integrantes:
        row = table.add_row().cells
        row[0].text = nome
        row[1].text = papel
        row[2].text = tempo
        row[3].text = slides
        row[4].text = desc
        for i, c in enumerate(row):
            set_cell_background(c, "F5F8FC" if len(table.rows)%2==0 else "FFFFFF")
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0:
                        r.font.bold = True
                        
    add_styled_heading(doc, "✅ Checklist de Sucesso para a Apresentação", level=2)
    checklist = [
        ("Postura e Oratória:", "Manter contato visual com a banca e falar de forma pausada e firme."),
        ("Transições Impecáveis:", "Cada integrante deve passar a palavra nominalmente para o próximo colega, demonstrando ensaio e coesão."),
        ("Domínio das Métricas:", "Ter na ponta da língua os números principais: Fator de Mérito, Orçamento de 1.8 e o p-valor da ANOVA (p < 0.05)."),
        ("Uso do Material Visual:", "Apontar para o GIF do campeão e para o Heatmap do MAP-Elites durante as explicações correspondentes.")
    ]
    for k, v in checklist:
        p = doc.add_paragraph()
        format_paragraph(p)
        r1 = p.add_run(f"• {k} ")
        r1.bold = True
        r1.font.color.rgb = RGBColor(20, 50, 110)
        p.add_run(v)
        
    docx_path = os.path.join(OUTPUT_DIR, "Divisao_Tarefas_Apresentadores.docx")
    doc.save(docx_path)
    
    # PDF
    pdf = CustomPDF("Divisão de Tarefas")
    pdf.add_title_block("DIVISÃO DE TAREFAS & APRESENTADORES", "Seminário de Inteligência Artificial — UNISENAI 2026")
    pdf.section_heading("1. Linha do Tempo e Atribuições por Integrante")
    
    for nome, papel, tempo, slides, desc in integrantes:
        pdf.sub_heading(f"{nome} — {papel} ({tempo}) | {slides}")
        pdf.body_text(desc)
        
    pdf.section_heading("2. Checklist de Sucesso do Grupo")
    for k, v in checklist:
        pdf.bullet_item(k, v)
        
    pdf_path = os.path.join(OUTPUT_DIR, "Divisao_Tarefas_Apresentadores.pdf")
    pdf.output(pdf_path)

# =========================================================================
# 2. GERAR: ROTEIRO DE APRESENTAÇÃO CRONOMETRADO (15 MIN)
# =========================================================================
def gerar_roteiro_apresentacao():
    print("Gerando Roteiro de Apresentacao (DOCX & PDF)...")
    
    # DOCX
    doc = Document()
    add_styled_heading(doc, "⏱️ Roteiro Cronometrado de Apresentação (15 Minutos)", level=1)
    
    p = doc.add_paragraph()
    format_paragraph(p)
    p.add_run("Projeto Mirage: IA Evolutiva para NPCs Evasivos\nEstrutura de 15 Minutos com Scripts de Fala e Transições Entre Integrantes").italic = True
    
    blocos = [
        ("Bloco 1: Abertura e Contextualização (0:00 às 3:00)", "Murilo Lameira", "Slides 1 a 3",
         "Boa noite aos membros da banca examinadora e ao professor orientador Ricardo Martinez Vicentini. Hoje apresentamos o projeto Mirage: Inteligência Artificial e Algoritmos Genéticos para NPCs Evasivos.\n\nNa indústria de jogos eletrônicos, os inimigos costumam ser programados com lógicas determinísticas rígidas, como máquinas de estados ou árvores de comportamento. O problema é que o jogador humano identifica esses padrões rapidamente, explora as brechas e o combate se torna previsível e monótono.\n\nO Mirage foi desenvolvido para quebrar essa previsibilidade. Em vez de regras manuais, colocamos os NPCs em uma arena contínua de física 2D sob uma chuva implacável de projéteis e usamos a computação evolutiva para que a própria IA descubra como sobreviver, desviar e equilibrar atributos físicos de combate.",
         "Para explicar como a literatura científica aborda a diversidade comportamental e os operadores evolutivos que impedem a IA de convergir para um único padrão previsível, passo a palavra ao Leonardo."),
         
        ("Bloco 2: Fundamentação e Qualidade-Diversidade (3:00 às 6:30)", "Leonardo Retori", "Slides 4 a 6",
         "Boa noite a todos. Na literatura clássica de Algoritmos Genéticos, a busca por uma única solução ótima global costuma falhar em jogos porque cria um 'inimigo padrão' que o jogador decora. Baseado na pesquisa de Kirk e Scirea sobre MAP-Elites (Qualidade e Diversidade), nós implementamos uma abordagem que mantém múltiplos arquétipos fenotípicos vivos e competitivos ao mesmo tempo: Tanques de alta vida, NPCs Equilibrados e Ninjas Evasivos de alta velocidade.\n\nPara isso, catalogamos os marcos geracionais no catálogo SEC (Skilled Experience Catalogue), inspirado em Glavin & Madden, que permite o Ajuste Dinâmico de Dificuldade (DDA) no jogo. Nossos operadores utilizam Seleção por Torneio com k=3 para calibrar a pressão seletiva, Crossover Uniforme para recombinação estrutural e Mutação Gaussiana com ruído estocástico, além de suporte ao modo EDS de mutação pura.",
         "Agora o Henry Matheus vai detalhar a cinemática física da arena, a restrição de orçamento global e a formulação matemática da função de fitness."),
         
        ("Bloco 3: Engenharia Física, Genes e Fitness (6:30 às 10:00)", "Henry Matheus", "Slides 7 a 9",
         "Boa noite. O motor físico do Mirage opera com integração temporal contínua a delta t = 0.05 segundos. A evasão é fundamentada nos Steering Behaviors de Craig Reynolds e na cinemática do Ponto de Maior Aproximação (CPA) descrita por Lee. A cada frame, o NPC analisa os vetores relativos de todos os projéteis e prevê o instante futuro exato da colisão mais crítica, aplicando uma força de repulsão perpendicular que o afasta do perigo.\n\nCada indivíduo carrega 4 genes contínuos: Vida, Ataque, Cadência e Velocidade. Para evitar o chamado 'Reward Hacking' e a criação de Super-NPCs invencíveis, instituímos um Orçamento Global de Atributos restrito a 1.8 pontos normalizados. Isso força escolhas táticas reais.\n\nNossa Função de Fitness é multi-objetivo e conta com a calibração por Fator de Mérito: sobreviver no modo Difícil rende 15.0 pontos por segundo e 20.0 pontos por esquiva perfeita, recompensando de forma justa quem sobrevive no ambiente de maior hostilidade.",
         "O Murilo Romualdo vai apresentar as mecânicas avançadas da arena, os resultados experimentais, a validação estatística rigorosa e a demonstração prática do campeão."),
         
        ("Bloco 4: Arena Dinâmica, Validação e Conclusão (10:00 às 14:00)", "Murilo Romualdo", "Slides 10 a 14",
         "Boa noite. Enriquecemos a arena com 4 pilares físicos de cobertura que absorvem tiros e bloqueiam a movimentação, incentivando a IA a usar cobertura tática (Line of Sight). Além dos disparos normais, criamos padrões complexos de Bullet Hell: disparos em leque divergente de 3 projéteis e vórtices espirais rotativos contínuos.\n\nPara eliminar o ruído estocástico de avaliações aleatórias (Noisy Fitness), rastreamos a melhor aptidão histórica, gerando curvas de convergência monotônicas em degraus de evolução.\n\nExecutamos 32 baterias automatizadas em segundo plano no Octave. Para validar cientificamente as diferenças entre dificuldades, aplicamos a One-Way ANOVA com cálculo exato de p-valor por função beta incompleta: obtivemos F = 138.24 e p = 1.44 x 10^-15, comprovando com significância matemática extrema que as estratégias evoluídas para cada dificuldade são distintas e reprodutíveis.\n\nConcluímos demonstrando que a IA Evolutiva é viável, de baixo custo computacional e capaz de criar oponentes dinâmicos e desafiadores para a indústria de jogos.",
         "Agradecemos a atenção de todos e estamos abertos às perguntas do orientador e da banca.")
    ]
    
    for titulo, apresentador, slides, fala, transicao in blocos:
        add_styled_heading(doc, titulo, level=2)
        
        p = doc.add_paragraph()
        format_paragraph(p)
        r = p.add_run(f"Apresentador: {apresentador} | {slides}")
        r.bold = True
        r.font.color.rgb = RGBColor(20, 50, 110)
        
        add_styled_heading(doc, "🗣️ Roteiro de Fala Sugerido:", level=3)
        p_fala = doc.add_paragraph()
        format_paragraph(p_fala, space_before=2, space_after=6)
        p_fala.add_run(f'"{fala}"').italic = True
        
        add_styled_heading(doc, "🔄 Frase de Transição:", level=3)
        p_trans = doc.add_paragraph()
        format_paragraph(p_trans, space_before=2, space_after=10)
        p_trans.add_run(f'"{transicao}"').bold = True
        
    docx_path = os.path.join(OUTPUT_DIR, "Roteiro_Apresentacao_15min.docx")
    doc.save(docx_path)
    
    # PDF
    pdf = CustomPDF("Roteiro de Apresentação (15 min)")
    pdf.add_title_block("ROTEIRO CRONOMETRADO DE APRESENTAÇÃO", "Scripts de Fala, Transições e Linha do Tempo (15 Minutos)")
    
    for titulo, apresentador, slides, fala, transicao in blocos:
        pdf.section_heading(f"{titulo} — {apresentador}")
        pdf.sub_heading(f"Referência: {slides}")
        pdf.body_text(fala, bold_prefix="Fala:")
        pdf.body_text(transicao, bold_prefix="Transição:")
        pdf.ln(2)
        
    pdf_path = os.path.join(OUTPUT_DIR, "Roteiro_Apresentacao_15min.pdf")
    pdf.output(pdf_path)

# =========================================================================
# 3. GERAR: DOSSIÊ COMPLETO DO SEMINÁRIO & DEFESA
# =========================================================================
def gerar_dossie_completo():
    print("Gerando Dossie Completo (DOCX & PDF)...")
    
    doc = Document()
    add_styled_heading(doc, "🛡️ Dossiê Científico e Guia de Apresentação — Mirage", level=1)
    
    p = doc.add_paragraph()
    format_paragraph(p)
    p.add_run("Centro Universitário SENAI SP — UNISENAI | Disciplina: Inteligência Artificial\nOrientador: Me. Ricardo Martinez Vicentini\nEquipe: Murilo Lameira, Leonardo Retori, Henry Matheus, Murilo Romualdo").italic = True
    
    add_styled_heading(doc, "1. Handout Científico e Fundamentação Matemática", level=2)
    
    secoes_cientificas = [
        ("O Paradigma da Evasão Adaptativa:", "Substituição de máquinas de estados determinísticas por computação evolutiva, eliminando a memorização de padrões pelo jogador."),
        ("O Cromossomo Contínuo de 4 Genes:", "HP [10, 200], Ataque [5, 50], Cadência de Disparo [0.5, 3.0 Hz] e Velocidade de Movimento [1.0, 8.0 m/s]."),
        ("Orçamento Global de Atributos (Point-Buy Budget):", "Restrição matemática sum(u_i) <= 1.8 para u_i em [0, 1]. Força trade-offs táticos reais (velocistas vs. tanques), impedindo Super-NPCs."),
        ("Cinemática de Esquiva de Craig Reynolds & CPA:", "Cálculo vetorial preditivo do Ponto de Maior Aproximação t_cpa = -dot(pr, vr) / |vr|^2. Desvios acionados apenas quando o perigo é iminente."),
        ("Função de Fitness Multi-Objetivo com Fator de Mérito:", "Equação: Fitness = max(0.1, (w1*T_survival) + (w2*N_dodge) + (w3*D_inflicted) - (p1*N_collision) - (p2*D_taken)). No Difícil, w1=15.0 e w2=20.0 recompensam a sobrevivência heróica."),
        ("Qualidade-Diversidade (MAP-Elites):", "Matriz 3x3 de nichos fenotípicos (Lento/Médio/Rápido vs. Tank/Balanceado/Glass Cannon), garantindo que a IA mantenha diferentes classes vivas."),
        ("Arena Dinâmica com Coberturas e Bullet Hell:", "4 pilares cilíndricos simétricos (r = 1.3m) que bloqueiam projéteis, além de rajadas em leque (Shotgun) e vórtices espirais contínuos."),
        ("Validação Estatística com One-Way ANOVA:", "Comprovação com 32 baterias independentes. F = 138.24 e p = 1.44 x 10^-15 (rejeição categórica de H0), comprovando que o comportamento evoluído não é aleatório.")
    ]
    
    for tit, desc in secoes_cientificas:
        p = doc.add_paragraph()
        format_paragraph(p)
        r = p.add_run(f"• {tit} ")
        r.bold = True
        r.font.color.rgb = RGBColor(20, 50, 110)
        p.add_run(desc)
        
    add_styled_heading(doc, "2. Blueprint Visual dos 14 Slides de Apresentação", level=2)
    
    slides = [
        ("Slide 1", "Murilo Lameira", "Capa Oficial", "Logo UNISENAI, título do Mirage, equipe e orientador Me. Ricardo Martinez Vicentini."),
        ("Slide 2", "Murilo Lameira", "O Problema do Determinismo", "Limitações das FSMs clássicas e como o jogador explora a previsibilidade."),
        ("Slide 3", "Murilo Lameira", "A Proposta do Mirage", "Apresentação da arena física 2D e do conceito de IA Evolutiva adaptativa."),
        ("Slide 4", "Leonardo Retori", "Qualidade-Diversidade (MAP-Elites)", "Teoria de Kirk & Scirea e o Heatmap 3x3 dos nichos de combate (map_elites_heatmap.png)."),
        ("Slide 5", "Leonardo Retori", "Catálogo SEC & Ajuste de Dificuldade", "Marcos de evolução off-line para o Dynamic Difficulty Adjustment (Glavin & Madden)."),
        ("Slide 6", "Leonardo Retori", "Operadores Genéticos do Mirage", "Seleção por Torneio (k=3), Crossover Uniforme e Mutação Gaussiana / Modo EDS."),
        ("Slide 7", "Henry Matheus", "Cinemática de Esquiva (Reynolds)", "Vetor de aproximação CPA e força de evasão perpendicular aos projéteis."),
        ("Slide 8", "Henry Matheus", "O Cromossomo e o Orçamento Global", "Os 4 genes e a restrição de orçamento (B = 1.8) que impede os Super-NPCs."),
        ("Slide 9", "Henry Matheus", "Função de Fitness & Fator de Mérito", "Equação multi-objetivo e escalonamento justo: Difícil > Médio > Fácil."),
        ("Slide 10", "Murilo Romualdo", "Arena 2D, Pilares e Bullet Hell", "Demonstração animada (demonstracao_npc.gif) com pilares de cobertura e disparos azuis."),
        ("Slide 11", "Murilo Romualdo", "Combate ao Reward Hacking", "Estudo de caso do exploit do tanque parado e a solução com penalidades calibradas."),
        ("Slide 12", "Murilo Romualdo", "Convergência Monotônica (Noisy Fitness)", "Eliminação de ruídos estocásticos através do rastreamento de melhor histórico global."),
        ("Slide 13", "Murilo Romualdo", "Validação Estatística (ANOVA & Boxplots)", "Resultados matemáticos: ANOVA com F = 138.24 e p = 1.44e-15 com Boxplots comparativos."),
        ("Slide 14", "Murilo Romualdo", "Conclusões & Engenharia de Software", "Síntese dos resultados acadêmicos, código aberto no GitHub e roadmap futuro.")
    ]
    
    tbl_slides = doc.add_table(rows=1, cols=4)
    tbl_slides.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_s = tbl_slides.rows[0].cells
    hdr_s[0].text = "Slide"
    hdr_s[1].text = "Apresentador"
    hdr_s[2].text = "Título do Slide"
    hdr_s[3].text = "Elementos Visuais e Conteúdo"
    for c in hdr_s:
        set_cell_background(c, "204080")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9.5)
                
    for sl, apr, tit, cont in slides:
        row = tbl_slides.add_row().cells
        row[0].text = sl
        row[1].text = apr
        row[2].text = tit
        row[3].text = cont
        for i, c in enumerate(row):
            set_cell_background(c, "F5F8FC" if len(tbl_slides.rows)%2==0 else "FFFFFF")
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i in (0, 1):
                        r.font.bold = True
                        
    add_styled_heading(doc, "3. Guia de Defesa Contra a Banca (FAQ Estratégico)", level=2)
    
    faq = [
        ("Como vocês comprovam que os resultados não foram fruto de sorte nos tiros aleatórios?",
         "Executamos 32 baterias independentes em paralelo e submetemos os dados à One-Way ANOVA com cálculo de p-valor exato via função beta incompleta (betainc). Obtivemos F = 138.24 e p = 1.44 x 10^-15 << 0.05. Além disso, todos os testes t pareados apresentaram p < 10^-8 e d de Cohen superior a 4.0, rejeitando a hipótese nula com rigor matemático."),
         
        ("Por que usar Algoritmos Genéticos e não Reinforcement Learning (Q-Learning / PPO)?",
         "Deep RL requer redes neurais densas e milhões de iterações de treino com altíssimo custo computacional, incompatíveis com a execução em tempo real em jogos eletrônicos. O AG associado ao MAP-Elites opera de forma extremamente leve diretamente nos parâmetros contínuos de controle e gera um catálogo diversificado de classes táticas em poucos segundos."),
         
        ("Por que a taxa de crossover é alta (70 a 90%) e a de mutação é baixa (5%)?",
         "O crossover é o operador responsável pela exploração construtiva, combinando os chamados 'building blocks' genéticos de sucesso. A mutação atua apenas como salvaguarda estocástica para evitar convergência prematura. Mutações acima de 15% quebram as estratégias consolidadas, transformando a evolução em busca aleatória caótica."),
         
        ("O que impedia o NPC de evoluir o exploit de virar um 'Tanque Parado' no centro?",
         "Três mecanismos articulados: 1) O Orçamento Global (B = 1.8) que impede o NPC de ter vida máxima e ataque máximo simultaneamente; 2) As penalidades progressivas por colisão (p1 e p2); e 3) Os padrões de tiro em espiral contínua e leque que cobrem a arena, tornando a imobilidade fatal."),
         
        ("Qual a vantagem prática de usar os Pilares de Cobertura na arena?",
         "Os pilares introduzem oclusão de projéteis (Line of Sight). Isso permite que a física de esquiva de Reynolds gere comportamentos emergentes onde o NPC aproveita obstáculos naturais como escudo tático, enriquecendo a dinâmica cinemática.")
    ]
    
    for q, a in faq:
        p_q = doc.add_paragraph()
        format_paragraph(p_q, space_before=6, space_after=2)
        r_q = p_q.add_run(f"❓ Pergunta: {q}")
        r_q.bold = True
        r_q.font.color.rgb = RGBColor(160, 40, 30)
        
        p_a = doc.add_paragraph()
        format_paragraph(p_a, space_before=1, space_after=6)
        r_a = p_a.add_run("💡 Resposta de Ouro: ")
        r_a.bold = True
        r_a.font.color.rgb = RGBColor(20, 80, 40)
        p_a.add_run(a)
        
    docx_path = os.path.join(OUTPUT_DIR, "Dossie_Completo_Seminario.docx")
    doc.save(docx_path)
    
    # PDF
    pdf = CustomPDF("Dossiê Científico e Apresentação")
    pdf.add_title_block("DOSSIÊ CIENTÍFICO E GUIA DE APRESENTAÇÃO", "Mirage: Inteligência Artificial e Algoritmos Genéticos para NPCs Evasivos")
    
    pdf.section_heading("1. Handout Científico e Fundamentação Teórica")
    for tit, desc in secoes_cientificas:
        pdf.bullet_item(tit, desc)
        
    pdf.section_heading("2. Blueprint dos 14 Slides de Apresentação")
    for sl, apr, tit, cont in slides:
        pdf.sub_heading(f"{sl} — {tit} ({apr})")
        pdf.body_text(cont)
        
    pdf.section_heading("3. Guia de Defesa contra a Banca (FAQ)")
    for q, a in faq:
        pdf.sub_heading(f"P: {q}")
        pdf.body_text(a, bold_prefix="R:")
        
    pdf_path = os.path.join(OUTPUT_DIR, "Dossie_Completo_Seminario.pdf")
    pdf.output(pdf_path)

# =========================================================================
# EXECUÇÃO PRINCIPAL
# =========================================================================
if __name__ == "__main__":
    gerar_divisao_tarefas()
    gerar_roteiro_apresentacao()
    gerar_dossie_completo()
    print("\n>>> SUCESSO: Todos os arquivos DOCX e PDF foram gerados com sucesso! <<<")
