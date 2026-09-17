# -*- coding: utf-8 -*-
"""
Script para Gerar o Documento Word (.docx) e PDF do:
GUIA DEFINITIVO DE ESTUDOS E SABATINA (PROJETO MIRAGE - UNISENAI 2026)
Separado nominalmente por integrante e com foco aprofundado na física.
"""

import os
import sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from fpdf import FPDF

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_DIR = os.path.join(BASE_DIR, "documentos_seminario")
DOCS_DIR = os.path.join(BASE_DIR, "docs")
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(DOCS_DIR, exist_ok=True)

# =========================================================================
# HELPERS DE FORMATAÇÃO DOCX
# =========================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def format_paragraph(p, space_before=2, space_after=4, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 6))
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(16, 44, 98) # Azul escuro institucional
    elif level == 2:
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(28, 68, 132)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(45, 90, 160)
    return h

def add_callout_box(doc, title, text_lines, bg_hex="F2F5F9", border_hex="1E4B8C"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Borda esquerda grossa
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=2, space_after=2)
    r_tit = p.add_run(f"📌 {title}\n")
    r_tit.bold = True
    r_tit.font.size = Pt(10.5)
    r_tit.font.color.rgb = RGBColor(16, 44, 98)
    
    for line in text_lines:
        p_l = cell.add_paragraph()
        format_paragraph(p_l, space_before=1, space_after=2)
        r_l = p_l.add_run(line)
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = RGBColor(40, 40, 40)
        
    doc.add_paragraph() # Espaçador

def add_qa_block(doc, question, answer):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "FAFBFD")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="D0DAE5"/><w:left w:val="single" w:sz="24" w:space="0" w:color="B82020"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="D0DAE5"/><w:right w:val="single" w:sz="6" w:space="0" w:color="D0DAE5"/></w:tcBorders>')
    tcPr.append(borders)
    
    # Pergunta
    p_q = cell.paragraphs[0]
    format_paragraph(p_q, space_before=2, space_after=3)
    r_q_label = p_q.add_run("❓ Pergunta da Banca: ")
    r_q_label.bold = True
    r_q_label.font.color.rgb = RGBColor(170, 20, 20)
    r_q_label.font.size = Pt(9.5)
    
    r_q = p_q.add_run(question)
    r_q.bold = True
    r_q.font.color.rgb = RGBColor(20, 20, 20)
    r_q.font.size = Pt(9.5)
    
    # Resposta
    p_a = cell.add_paragraph()
    format_paragraph(p_a, space_before=2, space_after=2)
    r_a_label = p_a.add_run("💡 Resposta de Mestre: ")
    r_a_label.bold = True
    r_a_label.font.color.rgb = RGBColor(18, 90, 45)
    r_a_label.font.size = Pt(9.5)
    
    r_a = p_a.add_run(answer)
    r_a.font.size = Pt(9.0)
    r_a.font.color.rgb = RGBColor(40, 40, 40)
    
    doc.add_paragraph()

# =========================================================================
# GERAÇÃO DO DOCUMENTO DOCX
# =========================================================================
def gerar_guia_docx():
    print("Gerando Guia de Estudo e Sabatina em Word (.docx)...")
    doc = Document()
    
    # Configuração de Margens da Página (1 polegada / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Cabeçalho e Rodapé
        header = section.header
        hp = header.paragraphs[0]
        format_paragraph(hp, space_before=0, space_after=0)
        hrun = hp.add_run("PROJETO MIRAGE — UNISENAI 2026 | INTELIGÊNCIA ARTIFICIAL")
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        footer = section.footer
        fp = footer.paragraphs[0]
        format_paragraph(fp, space_before=0, space_after=0)
        frun = fp.add_run("Guia de Estudos, Física e Sabatina da Banca — Engenharia de Controle e Automação")
        frun.font.size = Pt(8.0)
        frun.font.color.rgb = RGBColor(140, 140, 140)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # TÍTULO PRINCIPAL
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(title_p, space_before=0, space_after=4)
    r_t = title_p.add_run("📖 GUIA DEFINITIVO DE ESTUDOS E SABATINA")
    r_t.bold = True
    r_t.font.size = Pt(20)
    r_t.font.color.rgb = RGBColor(16, 44, 98)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(sub_p, space_before=0, space_after=8)
    r_s = sub_p.add_run("Projeto Mirage: Mecânica Vetorial, Cinemática Evasiva, Computação Evolutiva e Defesa Individual perante a Banca")
    r_s.italic = True
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = RGBColor(60, 60, 60)
    
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(meta_p, space_before=0, space_after=14)
    r_m = meta_p.add_run("Autores: Leonardo Retori • Henry Matheus • Murilo Lameira • Murilo Romualdo\nOrientador: Me. Ricardo Martinez Vicentini | UNISENAI São Caetano do Sul (2026)")
    r_m.font.size = Pt(9.5)
    r_m.bold = True
    r_m.font.color.rgb = RGBColor(40, 70, 120)

    # VISÃO GERAL
    add_styled_heading(doc, "🌌 1. Visão Panorâmica da Arquitetura do Mirage", level=1)
    p_intro = doc.add_paragraph()
    format_paragraph(p_intro)
    p_intro.add_run("O Mirage é um simulador cinemático e de otimização comportamental em tempo contínuo discretizado (Δt = 0.05 s / 20 Hz física / 50 FPS gráfico). Nele, um agente NPC combate em uma arena bidimensional de [-20, 20] x [-20, 20] metros contra padrões hostis de Bullet Hell. A grande inovação científica do trabalho reside na integração entre a mecânica clássica vetorial (Steering Behaviors de Reynolds e predição analítica de CPA) e a computação evolutiva (Algoritmos Genéticos com MAP-Elites e Orçamento Global de Atributos).")

    # TABELA DOS INTEGRANTES
    add_styled_heading(doc, "👥 2. Quadro Resumo de Atribuições & Foco Físico por Integrante", level=1)
    
    quadro = [
        ("Murilo Lameira", "Apresentador 1\n(0:00 - 3:00)", "Abertura, Gancho, FSMs vs Sistemas Contínuos e Visão Geral",
         "• Discretização temporal (Δt = 0.05s / 20 Hz)\n• Falha do determinismo booleano em espaços contínuos R²\n• Dinâmica de Newton (m = 1.0 kg => a = F)"),
        ("Leonardo Retori", "Apresentador 2\n(3:00 - 6:30)", "MAP-Elites (Kirk & Scirea), Catálogo SEC (DDA) e Operadores",
         "• Orçamento Global (Σ u_i <= 1.8) como conservador de capacidade física\n• Nichos fenotípicos de mobilidade (< 4.5, 4.5-6.5, > 6.5 m/s)\n• Mutação estocástica contínua via Box-Muller e modo EDS"),
        ("Henry Matheus", "Apresentador 3\n(6:30 - 10:00)", "Cinemática de Esquiva (CPA & Reynolds), Genes e Fitness",
         "• Dedução analítica exata do CPA: t_cpa = -(p_r · v_r) / ||v_r||²\n• Força de condução de Reynolds: F_evade = v_desejada - v_npc\n• Degradação cinemática de precisão de tiro em alta velocidade"),
        ("Murilo Romualdo", "Apresentador 4\n(10:00 - 14:00)", "Arena Dinâmica, Pilares, Bullet Hell, Bhandari e ANOVA",
         "• Pilares rígidos: absorção de tiros e deslizamento elástico tangencial\n• Padrões balísticos: Shotgun Cone (±12.6°) e Vórtice Espiral (ω = 4.5 rad/s)\n• Validação estatística ANOVA One-Way (F = 138.24, p = 1.44e-15)")
    ]
    
    t_int = doc.add_table(rows=1, cols=4)
    t_int.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_int.autofit = False
    
    hdr = t_int.rows[0].cells
    hdr[0].width = Inches(1.3)
    hdr[1].width = Inches(1.1)
    hdr[2].width = Inches(1.9)
    hdr[3].width = Inches(2.2)
    hdr[0].text = "Integrante"
    hdr[1].text = "Bloco & Tempo"
    hdr[2].text = "Tópico do Seminário"
    hdr[3].text = "Conexão Direta com a FÍSICA"
    for c in hdr:
        set_cell_background(c, "162C62")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
                
    for nome, bloco, top, fis in quadro:
        row = t_int.add_row().cells
        row[0].width = Inches(1.3)
        row[1].width = Inches(1.1)
        row[2].width = Inches(1.9)
        row[3].width = Inches(2.2)
        row[0].text = nome
        row[1].text = bloco
        row[2].text = top
        row[3].text = fis
        for i, c in enumerate(row):
            set_cell_background(c, "F5F8FC" if len(t_int.rows)%2==0 else "FFFFFF")
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)
                    if i == 0:
                        r.font.bold = True

    doc.add_paragraph()

    # COMPÊNDIO MESTRE DE FÍSICA E CINEMÁTICA
    add_styled_heading(doc, "🧭 3. Compêndio Mestre de Física e Cinemática Vetorial", level=1)
    p_comp = doc.add_paragraph()
    format_paragraph(p_comp)
    p_comp.add_run("Esta seção estabelece os pilares analíticos de toda a simulação da arena 2D. Todos os 4 integrantes devem dominar com clareza estas formulações:")
    
    add_styled_heading(doc, "3.1 Integração Numérica Semi-Implícita (Euler-Cromer)", level=2)
    add_callout_box(doc, "Equações de Movimento Translacional Discretizado (Δt = 0.05s)", [
        "A cada frame k, a força total Σ F_total acelera a massa unitária do NPC (m = 1.0 kg):",
        "v_npc(k+1) = truncate( v_npc(k) + (F_total / m) * Δt,  v_max )",
        "p_npc(k+1) = p_npc(k) + v_npc(k+1) * Δt",
        "Por que Euler Semi-Implícito? Porque ele utiliza a velocidade NOVA v(k+1) para atualizar a posição p(k+1). É um método simplético que preserva o volume do espaço de fase e a estabilidade de energia, impedindo que colisões com pilares gerem oscilações infinitas."
    ])

    add_styled_heading(doc, "3.2 Dedução Matemática Rigorosa do CPA (Closest Point of Approach)", level=2)
    add_callout_box(doc, "Passo a Passo da Derivação do CPA (Lee, 2014)", [
        "1. Posição e velocidade relativas: p_r = p_proj - p_npc  |  v_r = v_proj - v_npc",
        "2. Posição relativa em função de t: r(t) = p_r + v_r * t",
        "3. Distância quadrática euclidiana: D(t)² = ||r(t)||² = (p_r + v_r*t) · (p_r + v_r*t)",
        "   D(t)² = ||p_r||² + 2(p_r · v_r)t + ||v_r||² t²",
        "4. Minimização por derivada de primeira ordem: d/dt [D(t)²] = 2(p_r · v_r) + 2||v_r||² t = 0",
        "   => t_cpa = - (p_r · v_r) / ||v_r||²",
        "5. Segunda derivada: d²/dt² [D(t)²] = 2||v_r||² > 0 (comprovando que é estritamente um mínimo global).",
        "6. Gatilho de Esquiva: Ativado se 0 < t_cpa < 1.5s  E  ||r(t_cpa)|| < 2.5m  E  p_r · v_r < 0."
    ])

    add_styled_heading(doc, "3.3 Dinâmica dos Steering Behaviors (Craig Reynolds, 1999)", level=2)
    p_rey = doc.add_paragraph()
    format_paragraph(p_rey)
    p_rey.add_run("No instante t_cpa, o agente projeta a posição de fuga d_evade = p_npc(t_cpa) - p_proj(t_cpa). Se houver singularidade frontal exata (||d_evade|| < 1e-6), o código rotaciona o vetor do projétil em 90°: d_evade = [-v_py, v_px]. A velocidade desejada aponta para a fuga na velocidade terminal: v_desejada = (d_evade / ||d_evade||) * v_max. A força de direcionamento de Reynolds é: F_evade = v_desejada - v_npc.")

    add_styled_heading(doc, "3.4 Mecânica de Contato com Pilares (Restrição e Deslizamento Tangencial)", level=2)
    p_pil = doc.add_paragraph()
    format_paragraph(p_pil)
    p_pil.add_run("A arena possui 4 pilares em (±8, ±8) m com raio R = 1.3 m. Se ||p_npc - p_pilar|| < 2.3 m (soma dos raios), o NPC é projetado para a borda externa na direção normal unitária n = d_p / ||d_p||. Se a velocidade normal for penetrante (v_n = v · n < 0), ela é cancelada: v_npc <- v_npc - v_n * n. A velocidade tangencial é preservada, permitindo deslizamento suave ao redor da cobertura (Occlusion Steering). Projéteis que tocam o pilar (dist < 1.6m) são absorvidos e destruídos.")

    doc.add_page_break()

    # MÓDULO 1: MURILO LAMEIRA
    add_styled_heading(doc, "🧑‍💼 4. Módulo Nominal: Murilo Lameira", level=1)
    p_ml = doc.add_paragraph()
    format_paragraph(p_ml)
    p_ml.add_run("Papel: Apresentador 1 (0:00 às 3:00) | Tópico: Abertura, Contexto do Problema e Visão Geral do Mirage\n").bold = True
    p_ml.add_run("Missão: Abrir o seminário com postura firme, explicar por que os NPCs clássicos baseados em FSMs e Behavior Trees falham por determinismo excessivo e introduzir a ruptura do Mirage: agentes contínuos que aprendem física e esquiva via Algoritmos Genéticos.")

    add_styled_heading(doc, "Fundamentos e Relação com a Física", level=2)
    p_ml_f = doc.add_paragraph()
    format_paragraph(p_ml_f)
    p_ml_f.add_run("• Ruptura com o Grid Discreto: Jogos reais não operam em tabuleiros de xadrez; a arena do Mirage opera em R² contínuo com física integrada a 20 Hz (Δt = 0.05s).\n• Leis do Movimento: O agente tem massa de 1.0 kg e inércia; para mudar de direção, necessita aplicar forças reais, sem teletransporte.")

    add_styled_heading(doc, "Sabatina da Banca (Perguntas & Respostas)", level=2)
    add_qa_block(doc, 
                 "Por que não adicionar ruído aleatório (rand) a uma Máquina de Estados Finitos (FSM) em vez de implementar um Algoritmo Genético complexo?",
                 "Adicionar aleatoriedade estocástica a uma FSM quebra o determinismo, mas não gera inteligência. Um NPC aleatório frequentemente toma decisões catastróficas, como correr na direção de tiros ou colidir contra paredes. O Algoritmo Genético do Mirage descobre padrões ótimos de atributos físicos (velocidade e durabilidade) para sobrevivência no espaço contínuo, adaptando-se às exigências balísticas da arena.")
    
    add_qa_block(doc,
                 "Qual é o impacto do passo temporal Δt = 0.05s na física da simulação?",
                 "O passo de 0.05 segundos (20 Hz) é o ponto de equilíbrio ideal: garante a estabilidade numérica do integrador de Euler semi-implícito ao mesmo tempo em que permite executar centenas de simulações em lote paralelo em modo headless em poucos segundos.")

    add_styled_heading(doc, "Cartão de Bolso (Fórmulas e Métricas Chave)", level=2)
    add_callout_box(doc, "Cola Rápida — Murilo Lameira", [
        "Dimensões da Arena: [-20, 20] x [-20, 20] metros (Fronteiras seguras em ±18m).",
        "Frequência da Física: 20 Hz (Δt = 0.05s) | Renderização gráfica: ~50 FPS.",
        "Massa do NPC: m = 1.0 kg => Força resultante é numericamente igual à aceleração.",
        "Conceito Chave: Superar o determinismo previsível de FSMs através de sistemas dinâmicos contínuos e AG."
    ])

    doc.add_page_break()

    # MÓDULO 2: LEONARDO RETORI
    add_styled_heading(doc, "🧬 5. Módulo Nominal: Leonardo Retori", level=1)
    p_lr = doc.add_paragraph()
    format_paragraph(p_lr)
    p_lr.add_run("Papel: Apresentador 2 (3:00 às 6:30) | Tópico: Teoria Evolutiva, MAP-Elites, SEC e Operadores\n").bold = True
    p_lr.add_run("Missão: Apresentar o estado da arte em computação evolutiva aplicada a jogos: MAP-Elites (Kirk & Scirea, 2020) para preservar diversidade comportamental, Skilled Experience Catalogue (Glavin & Madden, 2015) para DDA, e os operadores genéticos implementados no código Octave.")

    add_styled_heading(doc, "Fundamentos e Relação com a Física", level=2)
    p_lr_f = doc.add_paragraph()
    format_paragraph(p_lr_f)
    p_lr_f.add_run("• Orçamento Global de Atributos (Point-Buy Budget): A soma normalizada dos 4 genes é limitada em Σ u_i <= 1.8. O ganho de velocidade (energia cinética) exige perda proporcional de HP ou dano.\n• Matriz MAP-Elites 3x3: Organizada em Mobilidade Física (Lento < 4.5, Médio 4.5-6.5, Rápido > 6.5 m/s) vs Classe de Durabilidade (Tanker ratio>3, Balanceado, Glass Cannon ratio<1).\n• Mutação Gaussiana com Box-Muller: Perturba atributos contínuos respeitando os limites do simplex físico.")

    add_styled_heading(doc, "Sabatina da Banca (Perguntas & Respostas)", level=2)
    add_qa_block(doc,
                 "Por que utilizar Seleção por Torneio (k=3) em vez da clássica Roleta Proporcional de Aptidão?",
                 "A Roleta Proporcional sofre de convergência prematura quando surge um indivíduo dominante no início do treino, e perde pressão seletiva no fim quando todos os fitness são semelhantes. O Torneio com k=3 baseia-se no ranking ordinal, mantendo a pressão seletiva estável e imune a variações de escala numérica da função de aptidão.")

    add_qa_block(doc,
                 "Como o MAP-Elites impede que a população convirja para um único 'Super-Ninja' veloz?",
                 "O MAP-Elites desacopla a competição: em vez de um único elite global, ele cria 9 nichos independentes. Um indivíduo Tanker com HP 200 e velocidade 2.5 m/s não compete com um Ninja de velocidade 8.0 m/s; ele compete apenas pela liderança do nicho de baixa mobilidade e alta absorção.")

    add_styled_heading(doc, "Cartão de Bolso (Fórmulas e Métricas Chave)", level=2)
    add_callout_box(doc, "Cola Rápida — Leonardo Retori", [
        "MAP-Elites: Matriz 3x3 (Eixo Y: Mobilidade | Eixo X: Razão HP/Ataque).",
        "Orçamento Normalizado: u_i = (G_i - G_min) / (G_max - G_min)  |  Σ u_i <= 1.8 (45% do teto).",
        "SEC (DDA): Exportação de marcos em 20%, 50% e 100% de evolução para data/catalogo_sec.csv.",
        "Modo EDS: Mutação pura sem crossover (Cr = 0, Mr = 30%) para adaptação tática ultrarrápida."
    ])

    doc.add_page_break()

    # MÓDULO 3: HENRY MATHEUS
    add_styled_heading(doc, "🦾 6. Módulo Nominal: Henry Matheus", level=1)
    p_hm = doc.add_paragraph()
    format_paragraph(p_hm)
    p_hm.add_run("Papel: Apresentador 3 (6:30 às 10:00) | Tópico: Arquitetura Física, Genes e Equação de Fitness\n").bold = True
    p_hm.add_run("Missão: Demonstrar o rigor cinemático do sistema: dedução analítica do CPA, equações de força de Craig Reynolds, estrutura dos 4 genes contínuos e a formulação multi-objetivo da fitness com fator de mérito por dificuldade.")

    add_styled_heading(doc, "Fundamentos e Relação com a Física", level=2)
    p_hm_f = doc.add_paragraph()
    format_paragraph(p_hm_f)
    p_hm_f.add_run("• Equação do CPA: t_cpa = -(p_r · v_r) / ||v_r||². O sinal negativo anula o produto escalar obtuso dos corpos em aproximação, fornecendo um tempo positivo.\n• Força de Direcionamento: F_evade = v_desejada - v_npc, direcionada ao longo de d_evade.\n• Penalidade de Precisão Dinâmica: Atirar em movimento reduz a estabilidade da mira em até 30%: accuracy = max(0.5, 1.0 - (||v||/v_max)*0.3).")

    add_styled_heading(doc, "Sabatina da Banca (Perguntas & Respostas)", level=2)
    add_qa_block(doc,
                 "O que ocorre se um projétil vier em rota de colisão frontal perfeita contra o NPC? Como o código evita divisão por zero?",
                 "Se o projétil estiver exatamente no centro do NPC no instante t_cpa, ||d_evade|| se torna zero. Em src/calculate_evade_force.m, existe uma salvaguarda que detecta essa condição singular e atribui um vetor ortogonal girado em 90 graus: d_evade = [-v_py, v_px]. Isso garante aceleração lateral perpendicular de desvio imediato.")

    add_qa_block(doc,
                 "Por que os pesos da função de fitness variam de acordo com a dificuldade?",
                 "No modo Fácil, os tiros são lentos e raros, então w1=2 e w2=2 bastam. No modo Difícil (Bullet Hell implacável), desviar de projéteis a 13.5 m/s é uma proeza cinética de alto risco, justificando w1=15 e w2=20 (Fator de Mérito Heroico), direcionando a seleção para maximizar desvios e sobrevivência.")

    add_styled_heading(doc, "Cartão de Bolso (Fórmulas e Métricas Chave)", level=2)
    add_callout_box(doc, "Cola Rápida — Henry Matheus", [
        "Fórmula do CPA: t_cpa = - (p_r · v_r) / ||v_r||²  |  Gatilho: 0 < t_cpa < 1.5s & dist < 2.5m.",
        "Reynolds Evade: F_evade = (d_evade / ||d_evade||) * v_max - v_npc.",
        "Vetor de 4 Genes: [HP: 10-200, Attack: 5-50, AtkSpeed: 0.5-3.0 Hz, MovSpeed: 1.0-8.0 m/s].",
        "Fitness: max( 0.1,  w1*T_surv + w2*N_dodge + w3*D_inflict - p1*N_coll - p2*D_taken )."
    ])

    doc.add_page_break()

    # MÓDULO 4: MURILO ROMUALDO
    add_styled_heading(doc, "📊 7. Módulo Nominal: Murilo Romualdo", level=1)
    p_mr = doc.add_paragraph()
    format_paragraph(p_mr)
    p_mr.add_run("Papel: Apresentador 4 (10:00 às 14:00) | Tópico: Arena Dinâmica, Balística, ANOVA, Demonstração e Fechamento\n").bold = True
    p_mr.add_run("Missão: Encerrar o seminário com dados empíricos rigorosos: explicar o comportamento físico dos pilares e disparos, revelar o estudo de caso de Reward Hacking superado, detalhar a parada antecipada de Bhandari, comprovar a significância estatística pela ANOVA One-Way (p << 0.05) e exibir a demonstração animada.")

    add_styled_heading(doc, "Fundamentos e Relação com a Física", level=2)
    p_mr_f = doc.add_paragraph()
    format_paragraph(p_mr_f)
    p_mr_f.add_run("• Pilares de Absorção e Deslizamento Tangencial: 4 obstáculos em (±8, ±8) m com raio 1.3m que absorvem tiros e projetam o NPC elasticamente com v_n anulada.\n• Balística Composta: Salvas em leque (Shotgun Cone ±12.6°) e Vórtice Espiral Danmaku (ω = 4.5 rad/s, v = 9.5 m/s).\n• Matriz de Hostilidade Física: Fácil (9 m/s, σ=0.25) -> Médio (11.5 m/s) -> Difícil (13.5 m/s, σ=0.08 cirúrgico).")

    add_styled_heading(doc, "Sabatina da Banca (Perguntas & Respostas)", level=2)
    add_qa_block(doc,
                 "Como vocês comprovam formalmente que os resultados obtidos não foram mero fruto do acaso estocástico?",
                 "Executamos baterias paralelas independentes e submetemos os dados consolidados ao teste One-Way ANOVA em src/teste_estatistico_hipoteses.m. Obtivemos a estatística F = 138.24 e p-valor exato p = 1.44 x 10^-15 via função beta incompleta (betainc). Como p << 0.05, rejeitamos categoricamente a hipótese nula H0, provando com significância estatística extrema a separabilidade das classes.")

    add_qa_block(doc,
                 "Por que as curvas dos modos Médio e Difícil pararam por volta da geração 21/22 enquanto o Fácil foi até 50?",
                 "Isso comprova a eficácia do Critério de Parada Antecipada de Bhandari (K=15 gerações, ε=1%). No Médio e Difícil, com elitismo e alto crossover (75-90%), a população estabilizou a melhor estratégia logo por volta da 6ª geração, permanecendo 15 gerações sem melhora > 1%, encerrando o processo para poupar CPU. No Fácil, a mutação altíssima de 15% sem elitismo manteve a variabilidade até o limite de 50 gerações.")

    add_styled_heading(doc, "Cartão de Bolso (Fórmulas e Métricas Chave)", level=2)
    add_callout_box(doc, "Cola Rápida — Murilo Romualdo", [
        "One-Way ANOVA: F = 138.24 | p = 1.44e-15 << 0.05 (Diferença fenotípica comprovada).",
        "Critério de Bhandari: Janela K = 15 gerações, limiar ε = 1% de melhoria.",
        "Pilares de Cobertura: 4 cilindros em (±8, ±8) m, raio 1.3m, absorção em dist < 1.6m.",
        "Vórtice Espiral Danmaku: r_vortex = [cos(4.5*t), sin(4.5*t)] * 9.5 m/s."
    ])

    doc.add_page_break()

    # GLOSSÁRIO MESTRE DE FÍSICA E VARIÁVEIS
    add_styled_heading(doc, "📖 8. Glossário Mestre de Variáveis, Constantes Físicas e Unidades SI", level=1)
    
    glossario = [
        ("Δt", "Passo de Tempo", "0.05", "s", "Discretização temporal do simulador (20 Hz física / 50 FPS)."),
        ("Ω", "Área da Arena", "[-20, 20]²", "m", "Plano cartesiano contínuo 2D de combate físico."),
        ("m", "Massa do Agente", "1.0", "kg", "Massa inercial unitária: F = m*a => a = F."),
        ("p_npc, v_npc", "Estado do NPC", "Contínuo", "m, m/s", "Posição e velocidade instantâneas do agente."),
        ("p_p, v_p", "Estado do Projétil", "Contínuo", "m, m/s", "Posição e velocidade instantâneas da bala."),
        ("p_r, v_r", "Vetores Relativos", "p_p - p_npc", "m, m/s", "Posição e velocidade relativas para cálculo do CPA."),
        ("t_cpa", "Tempo de Menor Distância", "-(p_r·v_r)/||v_r||²", "s", "Instante futuro projetado de máxima aproximação."),
        ("d_evade", "Vetor de Evasão", "p_npc - p_proj", "m", "Direção perpendicular de fuga calculada no instante t_cpa."),
        ("F_evade", "Força de Reynolds", "v_desejada - v_npc", "N", "Força de aceleração corretiva para condução cinemática."),
        ("r_hit", "Raio da Hitbox", "1.0", "m", "Envelope de colisão física direta (-25 HP de dano)."),
        ("R_radar", "Raio do Radar", "4.0", "m", "Envelope periférico de monitoramento de perigo (+1 Dodge)."),
        ("R_pillar", "Raio dos Pilares", "1.3", "m", "Obstáculos rígidos em (±8, ±8) m com absorção de balas."),
        ("B", "Orçamento Atributos", "1.8", "Adim.", "Teto global normalizado (Σ u_i <= 1.8, 45% do teto)."),
        ("G1 (HP)", "Pontos de Vida", "[10, 200]", "HP", "Integridade estrutural máxima do agente."),
        ("G2 (Attack)", "Poder de Ataque", "[5, 50]", "Dano", "Dano infligido por tiro de contra-ataque do NPC."),
        ("G3 (AtkSpeed)", "Cadência de Disparo", "[0.5, 3.0]", "Hz", "Frequência de contra-ataques emitidos pelo NPC."),
        ("G4 (MovSpeed)", "Velocidade Máxima", "[1.0, 8.0]", "m/s", "Velocidade terminal cinemática máxima (v_max)."),
        ("ω", "Velocidade Angular", "4.5", "rad/s", "Taxa de rotação do vórtice espiral no modo Difícil."),
        ("F", "Estatística ANOVA", "138.24", "Adim.", "Razão de variâncias entre grupos vs dentro de grupos."),
        ("p", "p-valor ANOVA", "1.44e-15", "Adim.", "Significância estatística extrema (p << 0.05).")
    ]
    
    t_gl = doc.add_table(rows=1, cols=5)
    t_gl.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_gl.autofit = False
    
    hdr_g = t_gl.rows[0].cells
    hdr_g[0].width = Inches(0.9)
    hdr_g[1].width = Inches(1.5)
    hdr_g[2].width = Inches(1.3)
    hdr_g[3].width = Inches(0.7)
    hdr_g[4].width = Inches(2.1)
    hdr_g[0].text = "Símbolo"
    hdr_g[1].text = "Nome Técnico"
    hdr_g[2].text = "Valor / Fórmula"
    hdr_g[3].text = "Unidade"
    hdr_g[4].text = "Função no Simulador"
    for c in hdr_g:
        set_cell_background(c, "162C62")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8.5)
                
    for s, n, v, u, d in glossario:
        row = t_gl.add_row().cells
        row[0].width = Inches(0.9)
        row[1].width = Inches(1.5)
        row[2].width = Inches(1.3)
        row[3].width = Inches(0.7)
        row[4].width = Inches(2.1)
        row[0].text = s
        row[1].text = n
        row[2].text = v
        row[3].text = u
        row[4].text = d
        for i, c in enumerate(row):
            set_cell_background(c, "F5F8FC" if len(t_gl.rows)%2==0 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    if i == 0:
                        r.font.bold = True

    # Salva o arquivo DOCX em ambos os diretórios
    docx_path1 = os.path.join(OUTPUT_DIR, "Guia_Estudo_e_Sabatina_Por_Integrante.docx")
    docx_path2 = os.path.join(DOCS_DIR, "Guia_Estudo_e_Sabatina_Por_Integrante.docx")
    doc.save(docx_path1)
    doc.save(docx_path2)
    print(f"DOCX salvo em: {docx_path1}")
    print(f"DOCX salvo em: {docx_path2}")

# =========================================================================
# GERAÇÃO DO DOCUMENTO PDF COM FPDF2
# =========================================================================
class StudyGuidePDF(FPDF):
    def __init__(self):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.set_auto_page_break(auto=True, margin=15)
        self.add_font("Arial", "", r"C:\Windows\Fonts\arial.ttf")
        self.add_font("Arial", "B", r"C:\Windows\Fonts\arialbd.ttf")
        self.add_font("Arial", "I", r"C:\Windows\Fonts\ariali.ttf")

    def header(self):
        self.set_font("Arial", "B", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 6, "PROJETO MIRAGE - UNISENAI 2026 | INTELIGÊNCIA ARTIFICIAL", border=0, align="L")
        self.cell(0, 6, "Guia de Estudos & Sabatina da Banca", border=0, align="R")
        self.ln(7)
        self.set_draw_color(210, 210, 210)
        self.line(10, 15, 200, 15)
        self.ln(3)

    def footer(self):
        self.set_y(-12)
        self.set_font("Arial", "I", 8)
        self.set_text_color(120, 120, 120)
        self.cell(0, 6, f"Página {self.page_no()}/{{nb}}", align="C")

    def title_block(self):
        self.add_page()
        self.set_font("Arial", "B", 16)
        self.set_text_color(16, 44, 98)
        self.multi_cell(0, 7, "GUIA DEFINITIVO DE ESTUDOS E SABATINA", align="C")
        self.ln(2)
        self.set_font("Arial", "I", 10)
        self.set_text_color(60, 60, 60)
        self.multi_cell(0, 5, "Mecânica Vetorial, Cinemática Evasiva, Computação Evolutiva e Defesa Individual perante a Banca", align="C")
        self.ln(2)
        self.set_font("Arial", "B", 8.5)
        self.set_text_color(40, 70, 120)
        self.multi_cell(0, 5, "Autores: Leonardo Retori • Henry Matheus • Murilo Lameira • Murilo Romualdo\nOrientador: Me. Ricardo Martinez Vicentini | UNISENAI São Caetano do Sul (2026)", align="C")
        self.ln(4)
        self.set_draw_color(24, 60, 120)
        self.set_line_width(0.6)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(5)

    def sec_heading(self, text):
        self.ln(3)
        self.set_font("Arial", "B", 12)
        self.set_text_color(16, 44, 98)
        self.cell(0, 6, text)
        self.ln(6)

    def sub_heading(self, text):
        self.ln(2)
        self.set_font("Arial", "B", 10)
        self.set_text_color(30, 70, 130)
        self.cell(0, 5, text)
        self.ln(5)

    def body_p(self, text):
        self.set_font("Arial", "", 9)
        self.set_text_color(30, 30, 30)
        self.multi_cell(0, 4.5, text)
        self.ln(2)

    def callout(self, title, lines):
        self.set_fill_color(242, 245, 250)
        self.set_draw_color(30, 75, 140)
        self.set_line_width(0.4)
        box_h = 7 + len(lines) * 4.5
        self.rect(10, self.get_y(), 190, box_h, style="FD")
        self.set_x(13)
        self.set_font("Arial", "B", 9)
        self.set_text_color(16, 44, 98)
        self.cell(0, 5, f"> {title}")
        self.ln(5)
        for l in lines:
            self.set_x(13)
            self.set_font("Arial", "", 8.5)
            self.set_text_color(40, 40, 40)
            self.cell(0, 4.2, l)
            self.ln(4.2)
        self.ln(4)

    def qa_item(self, q, a):
        self.set_font("Arial", "B", 9)
        self.set_text_color(160, 20, 20)
        self.write(4.5, "P: ")
        self.set_text_color(20, 20, 20)
        self.write(4.5, q + "\n")
        self.set_font("Arial", "B", 8.5)
        self.set_text_color(20, 90, 40)
        self.write(4.2, "R: ")
        self.set_font("Arial", "", 8.5)
        self.set_text_color(40, 40, 40)
        self.write(4.2, a + "\n\n")

def gerar_guia_pdf():
    print("Gerando Guia de Estudo e Sabatina em PDF...")
    pdf = StudyGuidePDF()
    pdf.title_block()
    
    pdf.sec_heading("1. Visão Panorâmica da Arquitetura do Mirage")
    pdf.body_p("O Mirage é um simulador cinemático e de otimização comportamental em tempo contínuo discretizado (Delta_t = 0.05 s / 20 Hz / 50 FPS). Nele, um agente NPC combate em uma arena 2D de [-20, 20] x [-20, 20] m contra padrões hostis de Bullet Hell. O trabalho integra mecânica clássica vetorial (Steering de Reynolds e CPA) com computação evolutiva (MAP-Elites e Orçamento Global de Atributos).")
    
    pdf.sec_heading("2. Compêndio Mestre de Física e Cinemática Vetorial")
    pdf.callout("Equações Fundamentais da Simulação", [
        "Euler Semi-Implícito: v(k+1) = truncate(v(k) + a*Delta_t, vmax)  |  p(k+1) = p(k) + v(k+1)*Delta_t",
        "Dedução do CPA: t_cpa = -(p_r · v_r) / ||v_r||²  (Derivada do produto escalar Euclidiano)",
        "Gatilho de Esquiva: 0 < t_cpa < 1.5s  &  dist_futura < 2.5m  &  p_r · v_r < 0",
        "Steering de Reynolds: F_evade = (d_evade / ||d_evade||)*v_max - v_npc",
        "Colisão com Pilares (R = 1.3m em +/-8, +/-8m): v_n cancelada se penetrante, v_tangencial mantida."
    ])
    
    # Módulos Nominais
    modulos = [
        ("3. Módulo Murilo Lameira (Abertura, Gancho & Visão Geral)", [
            ("Por que não adicionar ruído aleatório numa FSM em vez de usar Algoritmo Genético?",
             "Ruído aleatório torna o NPC imprevisível, mas não inteligente. Ele frequentemente comete suicídio contra tiros ou paredes. O AG descobre atributos físicos ideais para sobrevivência contínua."),
            ("Qual o impacto do passo de tempo Δt = 0.05s?",
             "O passo de 0.05s (20 Hz) garante a estabilidade do integrador simplético de Euler sem onerar a CPU em baterias de treino paralelas.")
        ]),
        ("4. Módulo Leonardo Retori (MAP-Elites, SEC & Operadores)", [
            ("Por que Seleção por Torneio (k=3) e não Roleta Proporcional?",
             "O Torneio depende apenas do ranking ordinal, eliminando o risco de convergência prematura por indivíduos super-dominantes da roleta."),
            ("Como o MAP-Elites impede que todos os NPCs fiquem iguais?",
             "Ele preserva o melhor indivíduo de cada um dos 9 nichos de mobilidade e classe de durabilidade, garantindo diversidade tática permanente.")
        ]),
        ("5. Módulo Henry Matheus (Física de Esquiva, Genes & Fitness)", [
            ("Mostre a dedução do CPA e o porquê do sinal negativo.",
             "Derivando D(t)² = ||p_r + v_r*t||² e igualando a zero, isola-se t_cpa = -(p_r·v_r)/||v_r||². O sinal negativo compensa o produto escalar negativo dos corpos que se aproximam."),
            ("O que ocorre em rota de colisão frontal perfeita?",
             "Se ||d_evade|| < 1e-6, o código aplica rotação de 90° no vetor do projétil [-v_py, v_px], forçando esquiva lateral imediata.")
        ]),
        ("6. Módulo Murilo Romualdo (Arena, ANOVA & Fechamento)", [
            ("Como provar estatisticamente que o aprendizado não foi sorte?",
             "Executamos One-Way ANOVA nos dados consolidando F = 138.24 e p = 1.44 x 10^-15 via betainc. Como p << 0.05, rejeita-se categoricamente a hipótese nula H0."),
            ("Por que o Médio e Difícil pararam na geração 21/22 e o Fácil foi até 50?",
             "Critério de Bhandari: estagnação de 15 gerações com melhora < 1%. No Difícil e Médio o elitismo convergiu rápido; no Fácil a mutação alta (15%) manteve o AG explorando até a geração 50.")
        ])
    ]
    
    for tit, qas in modulos:
        pdf.sec_heading(tit)
        for q, a in qas:
            pdf.qa_item(q, a)
            
    pdf_path1 = os.path.join(OUTPUT_DIR, "Guia_Estudo_e_Sabatina_Por_Integrante.pdf")
    pdf_path2 = os.path.join(DOCS_DIR, "Guia_Estudo_e_Sabatina_Por_Integrante.pdf")
    pdf.output(pdf_path1)
    pdf.output(pdf_path2)
    print(f"PDF salvo em: {pdf_path1}")
    print(f"PDF salvo em: {pdf_path2}")

if __name__ == "__main__":
    gerar_guia_docx()
    gerar_guia_pdf()
    print("\n>>> SUCESSO: Guia de Estudo e Sabatina gerado em DOCX e PDF! <<<")
